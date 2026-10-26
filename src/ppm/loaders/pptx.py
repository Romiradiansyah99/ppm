"""PPTX loader (python-pptx). Each slide becomes a text block; tables become
ParsedTables. Slide numbers are the section reference."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation

from ppm.loaders import ParsedTable, ParsedText


def parse_pptx(path: Path) -> tuple[list[ParsedTable], list[ParsedText]]:
    presentation = Presentation(str(path))
    tables: list[ParsedTable] = []
    texts: list[ParsedText] = []

    for slide_index, slide in enumerate(presentation.slides, start=1):
        lines: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in paragraph.runs).strip()
                    if text:
                        lines.append(text)
            if getattr(shape, "has_table", False):
                grid = [[cell.text.strip() for cell in row.cells] for row in shape.table.rows]
                grid = [row for row in grid if any(row)]
                if not grid:
                    continue
                tables.append(ParsedTable(
                    name=f"slide {slide_index} table {len(tables) + 1}",
                    start_row=2,
                    header=grid[0],
                    rows=grid[1:],
                ))
        if lines:
            texts.append(ParsedText(name=f"slide {slide_index}", text="\n".join(lines)))

    return tables, texts
