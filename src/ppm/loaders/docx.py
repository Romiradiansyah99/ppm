"""DOCX loader (python-docx). Paragraphs are grouped by heading so chunks cite
a section name, and tables keep their exact position in the document flow."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from ppm.loaders import ParsedTable, ParsedText


def _is_heading(paragraph: Paragraph) -> bool:
    style = paragraph.style.name if paragraph.style is not None else ""
    return style.lower().startswith(("heading", "title"))


def parse_docx(path: Path) -> tuple[list[ParsedTable], list[ParsedText]]:
    document = Document(str(path))
    tables: list[ParsedTable] = []
    texts: list[ParsedText] = []

    current_heading = "body"
    buffer: list[str] = []
    table_index = 0

    def flush() -> None:
        nonlocal buffer
        cleaned = [line for line in buffer if line.strip()]
        if cleaned:
            texts.append(ParsedText(name=current_heading, text="\n".join(cleaned).strip()))
        buffer = []

    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            paragraph = Paragraph(child, document)
            text = paragraph.text.strip()
            if not text:
                continue
            if _is_heading(paragraph):
                flush()
                current_heading = text
            else:
                buffer.append(text)
        elif child.tag == qn("w:tbl"):
            flush()
            table = Table(child, document)
            grid = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            grid = [row for row in grid if any(row)]
            if not grid:
                continue
            table_index += 1
            tables.append(ParsedTable(
                name=f"{current_heading} table {table_index}",
                start_row=2,
                header=grid[0],
                rows=grid[1:],
            ))

    flush()

    if not texts and not tables:
        # Empty document bodies still carry header/footer text worth indexing.
        for section in document.sections:
            header = "\n".join(p.text for p in section.header.paragraphs if p.text.strip())
            footer = "\n".join(p.text for p in section.footer.paragraphs if p.text.strip())
            if header or footer:
                texts.append(ParsedText(name="header/footer", text="\n".join(filter(None, [header, footer]))))

    return tables, texts
