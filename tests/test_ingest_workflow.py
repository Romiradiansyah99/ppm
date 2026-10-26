"""Ingest review/resume/dedupe without LangGraph."""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from ppm import db
from ppm.corpus.manifest import CorpusEntry
from ppm.models import reset_test_doubles, set_test_doubles
from ppm.workflows.ingest import IngestWorkflow
from tests.fakes import HashEmbedder

pytestmark = pytest.mark.db


def _docx(tmp_path: Path, text: str) -> Path:
    path = tmp_path / "sample.docx"
    document = Document()
    document.add_heading("Risk management", level=1)
    document.add_paragraph(text)
    document.save(path)
    return path


def _entry(path: Path) -> CorpusEntry:
    return CorpusEntry(
        path=path,
        relative_path="test/" + path.name,
        collection="test",
        licence_class="office_authored",
        index=True,
        is_template=True,
        stage="governance",
        category="risk",
        code="999",
        title="Risk test template",
        format="docx",
        is_blank=True,
        has_sample_content=False,
    )


@pytest.fixture(autouse=True)
def _fake_embedder():
    set_test_doubles(embedder=HashEmbedder())
    yield
    reset_test_doubles()


def test_review_resume_and_dedupe(tmp_path, settings, clean_db):
    workflow = IngestWorkflow(settings)
    entry = _entry(_docx(tmp_path, "Prepared for Barclays by the office."))
    result = workflow.run(entry)
    assert result["status"] == "awaiting_review"
    thread_id = result["thread_id"]

    row = db.fetch_one("SELECT stats FROM ingest_run WHERE thread_id = %s", (thread_id,))
    payload = row["stats"]["review"]
    assert "client_name_detected" in payload["reasons"]
    assert payload["redaction_hits"][0]["term"] == "Barclays"

    resumed = workflow.resume(thread_id, {
        "action": "approve_redacted",
        "licence_class": "office_authored",
        "confidentiality_class": "internal",
    })
    assert resumed["status"] in {"indexed", "indexed_text_only"}
    assert db.fetch_one("SELECT count(*) AS n FROM document_chunk")["n"] > 0
    assert db.fetch_one(
        "SELECT status FROM ingest_run WHERE thread_id = %s", (thread_id,)
    )["status"] == "completed"
    assert db.fetch_one(
        "SELECT ingest_status FROM document"
    )["ingest_status"] in {"indexed", "indexed_text_only"}

    second = workflow.run(entry)
    assert second["status"] == "duplicate"


def test_clean_document_completes_without_review(tmp_path, settings, clean_db):
    workflow = IngestWorkflow(settings)
    result = workflow.run(_entry(_docx(tmp_path, "A neutral project note.")))
    assert result["status"] in {"indexed", "indexed_text_only"}
    assert db.fetch_one("SELECT status FROM ingest_run")["status"] == "completed"


def test_reject_reference_only(tmp_path, settings, clean_db):
    workflow = IngestWorkflow(settings)
    result = workflow.run(_entry(_docx(tmp_path, "Prepared for Barclays by the office.")))
    assert result["status"] == "awaiting_review"
    resumed = workflow.resume(result["thread_id"], {"action": "forbid"})
    assert resumed["status"] == "forbidden"
    assert db.fetch_one("SELECT ingest_status FROM document")["ingest_status"] == "forbidden"
    assert db.fetch_one("SELECT count(*) AS n FROM ingest_staging")["n"] == 0
