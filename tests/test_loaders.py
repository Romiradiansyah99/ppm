from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

from ppm.loaders import UnsupportedFormatError, load_document


def test_docx_sections_and_tables(tmp_path: Path):
    path = tmp_path / "sample.docx"
    document = Document()
    document.add_heading("Scope", level=1)
    document.add_paragraph("This section describes cost control.")
    document.add_heading("Risk management", level=1)
    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "Risk"
    table.cell(0, 1).text = "Owner"
    table.cell(1, 0).text = "Delay"
    table.cell(1, 1).text = "PM"
    table.cell(2, 0).text = "Cost overrun"
    table.cell(2, 1).text = "QS"
    document.save(path)

    parsed = load_document(path)
    sections = {text.name for text in parsed.texts}
    assert "Scope" in sections
    assert len(parsed.tables) == 1
    assert parsed.tables[0].name.startswith("Risk management")
    assert parsed.tables[0].header == ["Risk", "Owner"]
    assert parsed.tables[0].rows[0] == ["Delay", "PM"]


def test_xlsx_header_detection(tmp_path: Path):
    path = tmp_path / "sample.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Rates"
    sheet.append(["Cost plan", None])
    sheet.append([None, None])
    sheet.append(["Code", "Description", "Rate"])
    sheet.append(["A1", "Concrete", 1000])
    workbook.save(path)

    parsed = load_document(path)
    table = parsed.tables[0]
    assert table.name == "Rates"
    assert table.header == ["Code", "Description", "Rate"]
    assert table.start_row == 4
    assert table.rows[0] == ["A1", "Concrete", 1000]


def test_pptx_slides(tmp_path: Path):
    path = tmp_path / "sample.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Cost report"
    box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1))
    box.text_frame.text = "Monthly cost position"
    presentation.save(path)

    parsed = load_document(path)
    assert parsed.texts
    assert parsed.texts[0].name == "slide 1"
    assert "Cost report" in parsed.texts[0].text


def test_legacy_doc_is_catalog_only(tmp_path: Path):
    path = tmp_path / "legacy.doc"
    path.write_bytes(b"not a real doc")
    with pytest.raises(UnsupportedFormatError):
        load_document(path)


def test_unknown_extension_is_rejected(tmp_path: Path):
    path = tmp_path / "archive.zip"
    path.write_bytes(b"zip")
    with pytest.raises(UnsupportedFormatError):
        load_document(path)
